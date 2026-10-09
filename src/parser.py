"""阶段1：文档解析。

把 data/raw/ 里的 PDF / DOCX / HTML / Markdown / TXT 统一转成
Document(text, metadata)。

Markdown / TXT 支持 YAML frontmatter（metadata 直接写在文件头）：

    ---
    title: 关于开展2026年研究生奖学金评审工作的通知
    department: 研究生院
    publish_date: 2026-09-15
    category: 奖学金
    year: 2026
    ---
    正文……

为什么 metadata 这么重要：后面按"只搜 2026 年研究生院的奖学金通知"
这类过滤，全靠它。真实场景里这些字段来自爬虫 / 数据库，第一版先手写在
frontmatter 里，结构完全一样。
"""
import os
import re
from pathlib import Path
from typing import Callable, List, Tuple

from src.schema import Document

try:
    import yaml
except ImportError:
    yaml = None

# 匹配文件开头的 YAML frontmatter（兼容 UTF-8 BOM）
FRONTMATTER_RE = re.compile(r"^﻿?\s*---\s*\n(.*?)\n---\s*\n?", re.DOTALL)


def _parse_frontmatter(text: str) -> Tuple[dict, str]:
    m = FRONTMATTER_RE.match(text)
    if not m:
        return {}, text
    raw, body = m.group(1), text[m.end():]
    if yaml is None:  # 没装 pyyaml 时的兜底：极简 key: value 解析
        meta = {}
        for line in raw.splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip()
        return meta, body
    try:
        meta = yaml.safe_load(raw) or {}
    except Exception:
        meta = {}
    return meta, body


def _load_sidecar_metadata(path: str) -> dict:
    """读取与源文件同名的 .meta.yaml，例如 handbook.pdf -> handbook.meta.yaml。"""
    sidecar = Path(path).with_suffix(".meta.yaml")
    if not sidecar.exists():
        return {}
    if yaml is None:
        raise RuntimeError("读取 metadata sidecar 需要 pyyaml 依赖。")
    try:
        payload = yaml.safe_load(sidecar.read_text(encoding="utf-8")) or {}
    except yaml.YAMLError as exc:
        raise ValueError(f"metadata sidecar 格式错误: {sidecar}") from exc
    if not isinstance(payload, dict):
        raise ValueError(f"metadata sidecar 必须是 YAML 对象: {sidecar}")
    return payload


def _text_file(path: str) -> Document:
    with open(path, encoding="utf-8") as f:
        text = f.read()
    meta, body = _parse_frontmatter(text)
    return Document(text=body.strip(), metadata=meta)


def _pdf_file(path: str) -> Document:
    import pymupdf
    with pymupdf.open(path) as doc:
        pages = [page.get_text() for page in doc]
        native_text = "\n".join(pages).strip()
        visible_chars = len(re.sub(r"\s+", "", native_text))
        minimum_chars = max(100, len(doc) * 20)
        if visible_chars >= minimum_chars:
            return Document(text=native_text, metadata={})

        try:
            pages = [
                page.get_text(textpage=page.get_textpage_ocr(language="chi_sim+eng", dpi=200, full=True))
                for page in doc
            ]
        except RuntimeError as exc:
            raise RuntimeError(
                f"{path} 几乎没有可提取文字，且 OCR 启动失败。"
                "请安装 Tesseract OCR 及 chi_sim 中文语言包后重试。"
            ) from exc
        text = "\n".join(pages)
    return Document(text=text.strip(), metadata={})


def _docx_file(path: str) -> Document:
    import docx
    d = docx.Document(path)
    text = "\n".join(p.text for p in d.paragraphs)
    return Document(text=text.strip(), metadata={})


def _html_file(path: str) -> Document:
    from bs4 import BeautifulSoup
    with open(path, encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "lxml")
    return Document(text=soup.get_text("\n").strip(), metadata={})


HANDLERS = {
    ".md": _text_file, ".txt": _text_file, ".markdown": _text_file,
    ".pdf": _pdf_file, ".docx": _docx_file,
    ".html": _html_file, ".htm": _html_file,
}


def load_file(path: str) -> Document:
    ext = os.path.splitext(path)[1].lower()
    handler = HANDLERS.get(ext)
    if handler is None:
        raise ValueError(f"不支持的文件类型: {ext}")
    try:
        doc = handler(path)
    except ImportError as e:
        raise RuntimeError(
            f"解析 {path} 需要额外依赖，请先安装：{e.name}"
        ) from e
    doc.metadata.update(_load_sidecar_metadata(path))
    # 来源字段由系统维护，不能被 frontmatter 或 sidecar 覆盖。
    doc.metadata["source_path"] = path
    doc.metadata["format"] = ext.lstrip(".")
    doc.metadata["document_id"] = os.path.basename(path)
    return doc


def load_directory(
    root: str,
    on_error: Callable[[str, Exception], None] | None = None,
) -> List[Document]:
    """递归加载支持的文件；提供 on_error 时跳过失败文件并报告错误。"""
    docs: List[Document] = []
    for dirpath, _, files in os.walk(root):
        for fn in sorted(files):
            if os.path.splitext(fn)[1].lower() in HANDLERS:
                path = os.path.join(dirpath, fn)
                try:
                    docs.append(load_file(path))
                except Exception as exc:
                    if on_error is None:
                        raise
                    on_error(path, exc)
    return docs
